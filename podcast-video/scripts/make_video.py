"""把已有播客音频和幻灯片 PNG 合成视频，并导出 PPTX。
时间轴：每句 = segments/NNN_X.wav 时长 + GAP（与 make_podcast.py 拼接方式一致），
所以切页时间和 podcast.mp3 精确对齐。"""
import importlib.util, subprocess, sys
from pathlib import Path

ROOT = Path(sys.argv[1]).resolve()  # 用法：make_video.py <章节目录>
V = ROOT / "video"
GAP = 0.35
spec = importlib.util.spec_from_file_location("spec", V / "spec.py")
mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
SLIDES, SLUG = mod.SLIDES, mod.SLUG

def dur(p):
    return float(subprocess.check_output(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                                          "-of", "csv=p=0", str(p)]).strip())

wavs = sorted((ROOT / "segments").glob("[0-9][0-9][0-9]_?.wav"))
starts, t = [], 0.0
for w in wavs:
    starts.append(t); t += dur(w) + GAP
lens = [dur(w) for w in wavs]
total = t
texts = [l for l in (ROOT / "script.txt").read_text(encoding="utf-8").splitlines() if l[:2] in ("A:", "B:")]
assert len(texts) == len(wavs), (len(texts), len(wavs))

# 3) 画面：每页持续到下一页第一句开始，按时长拼接幻灯片（不加字幕）
lst = V / "build" / "frames.txt"
with lst.open("w") as f:
    for k, s in enumerate(SLIDES):
        a = starts[s["start"]]
        z = starts[SLIDES[k + 1]["start"]] if k + 1 < len(SLIDES) else total
        f.write(f"file '{V}/build/png/slide{k+1:02d}.png'\nduration {z - a:.3f}\n")
    f.write(f"file '{V}/build/png/slide{len(SLIDES):02d}.png'\n")

subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(lst),
                "-i", str(ROOT / "podcast.mp3"),
                "-vf", "fps=10,format=yuv420p", "-map", "0:v", "-map", "1:a",
                "-c:v", "libx264", "-preset", "veryfast", "-tune", "stillimage", "-crf", "20", "-g", "100",
                "-c:a", "aac", "-b:a", "160k", "-shortest", "-movflags", "+faststart",
                str(V / f"{SLUG}.mp4")], check=True)

# 4) PPTX：每页整幅幻灯片图 + 演讲者备注为对应的对话
from pptx import Presentation
from pptx.util import Emu
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
for k, s in enumerate(SLIDES):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.shapes.add_picture(str(V / "build" / "png" / f"slide{k+1:02d}.png"), 0, 0,
                          prs.slide_width, prs.slide_height)
    sl.notes_slide.notes_text_frame.text = s["notes"]
prs.save(str(V / f"{SLUG}.pptx"))
print(f"done: {len(SLIDES)} slides, {total:.1f}s")
